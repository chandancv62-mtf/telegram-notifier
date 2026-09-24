import os
import datetime
import pandas as pd
import yfinance as yf
from tqdm import tqdm
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

def apply_excel_styling(ws):
    center_align = Alignment(horizontal='center', vertical='center')

    font_header = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    fill_header = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')

    font_red = Font(name='Calibri', size=11, color='FF0000', bold=True)
    font_green = Font(name='Calibri', size=11, color='008000', bold=True)
    font_default = Font(name='Calibri', size=11, color='000000')

    fill_target = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
    fill_sl = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
    fill_active = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
    fill_expired = PatternFill(start_color='D9D9D9', end_color='D9D9D9', fill_type='solid')

    header_map = {str(ws.cell(row=1, column=c).value): c for c in range(1, ws.max_column + 1)}

    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            cell.alignment = center_align

            if row == 1:
                cell.font = font_header
                cell.fill = fill_header
            else:
                col_name_list = [k for k, v in header_map.items() if v == col]
                if col_name_list:
                    col_name = col_name_list[0]
                    if col_name in ['SL', 'SL_%(Point)']:
                        cell.font = font_red
                    elif col_name in ['ENTRY', 'TARGET_1:1', 'TARGET_1:2', 'PnL_%']:
                        cell.font = font_green
                    else:
                        cell.font = font_default

                    if col_name == 'Trade_Status':
                        val = str(cell.value)
                        if 'TARGET' in val:
                            cell.fill = fill_target
                            cell.font = font_green
                        elif 'SL HIT' in val:
                            cell.fill = fill_sl
                            cell.font = font_red
                        elif 'ACTIVE' in val:
                            cell.fill = fill_active
                        elif 'EXPIRED' in val:
                            cell.fill = fill_expired
                            cell.font = Font(name='Calibri', size=11, color='595959', italic=True)

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 2, 8)


def track_trade_status(input_file='weekly_final_trading_signals.xlsx', output_file='trade_tracker_results.xlsx'):
    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_root = os.path.dirname(script_dir)
    
    # Path resolution for input file (check data/ dir, then root)
    input_path = os.path.join(project_root, 'data', os.path.basename(input_file))
    if not os.path.exists(input_path):
        input_path = os.path.join(project_root, os.path.basename(input_file))

    # Save output strictly inside data/ directory
    data_dir = os.path.join(project_root, 'data')
    os.makedirs(data_dir, exist_ok=True)
    output_path = os.path.join(data_dir, os.path.basename(output_file))

    df_valid = pd.DataFrame()

    try:
        if os.path.exists(input_path):
            df_valid = pd.read_excel(input_path, sheet_name='Valid_Setups_Only')
    except Exception as e:
        print(f"⚠️ Warning reading '{input_file}': {e}")

    today = datetime.datetime.now()
    monday_start = (today - datetime.timedelta(days=today.weekday())).replace(hour=0, minute=0, second=0, microsecond=0)
    sunday_end = monday_start + datetime.timedelta(days=6, hours=23, minutes=59)

    tracked_results = []

    if df_valid.empty:
        print("ℹ️ Valid setups sheet is empty for current week. Clearing 'current_week' sheet...")
    else:
        print(f"\n🚀 Tracking Trade Outcomes for Current Week ({len(df_valid)} Valid Stocks)...\n")

        for idx, row in tqdm(df_valid.iterrows(), total=len(df_valid), desc="Tracking Trades", unit="stock"):
            stock_name = str(row['Ticker']).strip()
            ticker_symbol = stock_name if (stock_name.endswith('.NS') or stock_name.endswith('.BO')) else stock_name + '.NS'

            sl_price = float(row['SL'])
            entry_price = float(row['ENTRY'])
            bu_time_str = str(row['BU_Time'])

            risk = round(entry_price - sl_price, 2)
            sl_pct = round((risk / entry_price) * 100, 2)
            sl_pct_with_points = f"{sl_pct}% ({risk})"

            target_1_1 = float(row['TARGET_1:1']) if 'TARGET_1:1' in row else round(entry_price + (1 * risk), 2)
            target_1_2 = float(row['TARGET_1:2']) if 'TARGET_1:2' in row else round(entry_price + (2 * risk), 2)

            trade_status = "WAIT (No Entry)"
            entry_triggered_time = "N/A"
            exit_time = "N/A"
            exit_price = None
            pnl_pct = 0.0

            try:
                ticker = yf.Ticker(ticker_symbol)
                df_15m = ticker.history(period="5d", interval="15m")

                if not df_15m.empty:
                    if df_15m.index.tz is not None:
                        df_15m.index = df_15m.index.tz_localize(None)

                    bu_dt = pd.to_datetime(bu_time_str)
                    df_after_bu = df_15m[(df_15m.index > bu_dt) & (df_15m.index >= monday_start) & (df_15m.index <= sunday_end)]
                    df_after_bu = df_after_bu.between_time('09:15', '15:15')

                    entry_candles = df_after_bu[df_after_bu['High'] >= entry_price]

                    if not entry_candles.empty:
                        first_entry_candle = entry_candles.iloc[0]
                        entry_triggered_time = first_entry_candle.name.strftime('%Y-%m-%d %H:%M')
                        trade_status = "ACTIVE"

                        df_after_entry = df_after_bu[df_after_bu.index >= first_entry_candle.name]

                        target_1_1_hit = False
                        last_candle_dt = None
                        last_candle_close = None

                        for dt, c_row in df_after_entry.iterrows():
                            candle_high = float(c_row['High'])
                            candle_low = float(c_row['Low'])
                            last_candle_dt = dt
                            last_candle_close = float(c_row['Close'])

                            if candle_low <= sl_price:
                                trade_status = "TARGET 1:1 HIT (Then SL)" if target_1_1_hit else "SL HIT"
                                exit_time = dt.strftime('%Y-%m-%d %H:%M')
                                exit_price = round(sl_price, 2)
                                break

                            if candle_high >= target_1_1 and not target_1_1_hit:
                                target_1_1_hit = True
                                trade_status = "TARGET 1:1 HIT (Active for 1:2)"
                                exit_time = dt.strftime('%Y-%m-%d %H:%M')
                                exit_price = round(target_1_1, 2)

                            if candle_high >= target_1_2:
                                trade_status = "TARGET 1:2 HIT"
                                exit_time = dt.strftime('%Y-%m-%d %H:%M')
                                exit_price = round(target_1_2, 2)
                                break

                        if exit_price is None and last_candle_dt is not None:
                            exit_time = last_candle_dt.strftime('%Y-%m-%d %H:%M') + " (Active)"
                            exit_price = round(last_candle_close, 2)

                        if exit_price is not None:
                            pnl_pct = round(((exit_price - entry_price) / entry_price) * 100, 2)

                    else:
                        trade_status = "WAIT (No Entry)"

            except Exception as e:
                trade_status = f"ERROR: {str(e)}"

            res = {
                'Ticker': stock_name,
                'PWL': row['PWL'],
                'SL': sl_price,
                'ENTRY': entry_price,
                'SL_%(Point)': sl_pct_with_points,
                'TARGET_1:1': target_1_1,
                'TARGET_1:2': target_1_2,
                'Entry_Triggered_Time': entry_triggered_time,
                'Trade_Status': trade_status,
                'Exit_Time': exit_time,
                'Exit_Price': exit_price,
                'PnL_%': pnl_pct
            }
            tracked_results.append(res)

    if tracked_results:
        df_current = pd.DataFrame(tracked_results)
    else:
        df_current = pd.DataFrame(columns=['Ticker', 'PWL', 'SL', 'ENTRY', 'SL_%(Point)', 'TARGET_1:1', 'TARGET_1:2', 'Entry_Triggered_Time', 'Trade_Status', 'Exit_Time', 'Exit_Price', 'PnL_%'])

    df_history = pd.DataFrame()

    if not df_current.empty:
        df_triggered_only = df_current[
            (df_current['Entry_Triggered_Time'].astype(str) != "N/A") & 
            (~df_current['Trade_Status'].astype(str).str.contains("EXPIRED", case=False, na=False))
        ].copy()
    else:
        df_triggered_only = pd.DataFrame()

    if os.path.exists(output_path):
        try:
            df_existing_history = pd.read_excel(output_path, sheet_name='history_data')
            df_history = pd.concat([df_existing_history, df_triggered_only], ignore_index=True)
            df_history.drop_duplicates(subset=['Ticker', 'Entry_Triggered_Time'], keep='last', inplace=True)
        except Exception:
            df_history = df_triggered_only.copy()
    else:
        df_history = df_triggered_only.copy()

    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        df_current.to_excel(writer, sheet_name='current_week', index=False)
        apply_excel_styling(writer.sheets['current_week'])

        if not df_history.empty:
            df_history.to_excel(writer, sheet_name='history_data', index=False)
            apply_excel_styling(writer.sheets['history_data'])

    print(f"\n✅ Tracking Complete in Data Folder!")
    print(f"📊 Sheet 'current_week': Reset & Updated ({len(df_current)} rows).")
    print(f"📚 Sheet 'history_data': Safely holding {len(df_history)} total executed trades.")

if __name__ == "__main__":
    track_trade_status()
